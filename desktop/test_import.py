import sys
from pathlib import Path

# The portable app packages these existing importers beside the desktop sources.
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from app.document_reader import extract_text


def test_utf8_bom(tmp_path):
    path=tmp_path/'text.txt';path.write_text('Первая строка.',encoding='utf-8-sig')
    assert extract_text(path)=='Первая строка.'


def test_docx_preserves_table_position(tmp_path):
    from docx import Document
    doc=Document();doc.add_paragraph('Первый абзац.')
    doc.add_table(rows=1,cols=1).cell(0,0).text='Таблица посередине.'
    doc.add_paragraph('Последний абзац.')
    path=tmp_path/'test.docx';doc.save(path)
    text=extract_text(path)
    assert text.index('Первый')<text.index('Таблица')<text.index('Последний')


def test_epub_preserves_spine_order(tmp_path):
    from ebooklib import epub
    book=epub.EpubBook();book.set_identifier('test');book.set_title('Книга');book.set_language('ru')
    first=epub.EpubHtml(title='Первая',file_name='first.xhtml',lang='ru')
    first.content='<p>Первая глава.</p>'
    second=epub.EpubHtml(title='Вторая',file_name='second.xhtml',lang='ru')
    second.content='<p>Вторая глава.</p>'
    book.add_item(second);book.add_item(first)
    book.add_item(epub.EpubNcx());book.add_item(epub.EpubNav())
    book.spine=[first,second]
    path=tmp_path/'test.epub';epub.write_epub(str(path),book)
    text=extract_text(path)
    assert text.index('Первая')<text.index('Вторая')


def test_structure_docx_headings_and_tables(tmp_path):
    from docx import Document
    from document_structure import extract_sections
    doc=Document();doc.add_heading('Часть первая',level=1);doc.add_paragraph('Вступительный текст.')
    doc.add_heading('Глава первая',level=2);doc.add_table(rows=1,cols=1).cell(0,0).text='Текст таблицы.'
    path=tmp_path/'structured.docx';doc.save(path)
    result=extract_sections(path)
    assert [(s['title'],s['level']) for s in result]==[('Часть первая',1),('Глава первая',2)]
    assert 'Текст таблицы.' in result[1]['text']


def test_structure_plain_and_html(tmp_path):
    from document_structure import extract_sections
    text=tmp_path/'book.txt';text.write_text('Предисловие.\nЧасть I\nТекст части.\nГлава 1\nТекст главы.',encoding='utf-8')
    result=extract_sections(text)
    assert len(result)==3 and result[-1]['level']==2
    html=tmp_path/'book.html';html.write_text('<h1>Часть <b>один</b></h1><p>Начало.</p><h2>Глава</h2><p>Конец.</p>',encoding='utf-8')
    result=extract_sections(html)
    assert len(result)==2 and result[0]['text'].count('Часть один')==1
    assert 'Конец.' in result[1]['text']
