import json
import threading
import time
from PySide6.QtCore import QTimer
from PySide6.QtTest import QTest
from test_ui import app, window
from project import Project, split_text
from test_project import FakeEngine
from test_ui import click_row
from PySide6.QtCore import Qt


def wait_until(app,predicate,seconds=30):
    deadline=time.monotonic()+seconds
    while not predicate() and time.monotonic()<deadline:
        app.processEvents();QTest.qWait(10)
    assert predicate(), 'Background operation timed out'


def empty(window):
    p=window.project
    p.remove_segments([s['id'] for _,s in p.rows(p.data['root'])]);window.refresh()
    return p


def test_2500_chapters_progressive_import_and_concurrent_generation(window,app,tmp_path,monkeypatch):
    p=empty(window)
    p.data['settings']['retry_low_score']=False  # This stress test measures import/UI, not retry policy.
    source=tmp_path/'large.md'
    source.write_text(''.join(f'# Глава {i}\n'+('Текст большой книги для проверки. '*35)+'\n' for i in range(2500)),encoding='utf-8')
    monkeypatch.setattr(window,'import_analyzed',lambda count:window.importer.choose(True))
    window.engine=FakeEngine(fail=True)
    ticks=[];timer=QTimer();timer.setInterval(10)
    timer.timeout.connect(lambda:ticks.append(time.monotonic()));timer.start()
    started=time.monotonic()
    window.start_import(source)
    wait_until(app,lambda:len(p.data['nodes'])>=2)
    assert window.import_active()
    chapter=p.data['nodes'][1]
    window.tree.setCurrentItem(window.tree_items[chapter['id']])
    window.generate()
    wait_until(app,lambda:not window.import_active() and not window.busy(),60)
    timer.stop()
    assert len(p.data['nodes'])==2501
    assert p.counts(chapter['id'])[1]>0
    assert len(ticks)>10
    max_gap=max(b-a for a,b in zip(ticks,ticks[1:]))
    assert max_gap<1.5
    print(f'STRESS: 2500 chapters, {len(p.rows(p.data["root"]))} fragments, '
          f'{time.monotonic()-started:.2f}s, UI timer max gap {max_gap:.3f}s, {len(ticks)} ticks')
    assert window.table.rowCount()<=300
    assert window.import_progress.value()==100
    p.flush_saves()
    saved=json.loads((p.folder/'project.json').read_text(encoding='utf-8'))
    assert len(saved['nodes'])==2501
    assert all(n['segments'] for n in saved['nodes'][1:])


def test_cancel_import_keeps_completed_batches(window,app,tmp_path,monkeypatch):
    p=empty(window);source=tmp_path/'cancel.md'
    source.write_text(''.join(f'# Глава {i}\nТекст.\n' for i in range(2500)),encoding='utf-8')
    monkeypatch.setattr(window,'import_analyzed',lambda count:window.importer.choose(True))
    window.start_import(source)
    wait_until(app,lambda:len(p.data['nodes'])>2)
    window.cancel_import()
    wait_until(app,lambda:not window.import_active())
    assert 2<len(p.data['nodes'])<2501
    p.flush_saves()
    assert len(json.loads((p.folder/'project.json').read_text(encoding='utf-8'))['nodes'])==len(p.data['nodes'])
    assert not window.busy()


def test_loading_runs_off_gui_thread_and_allows_generation(window,app,tmp_path,monkeypatch):
    import document_import
    main_thread=threading.get_ident();threads=[];release=threading.Event()
    def slow_blocks(path,report,stop):
        threads.append(threading.get_ident());report('Загрузка',1,10)
        while not release.wait(.01):
            if stop.is_set():return
        yield 'Импортированный текст.',None
    monkeypatch.setattr(document_import,'read_blocks',slow_blocks)
    window.engine=FakeEngine()
    window.start_import(tmp_path/'fake.txt')
    wait_until(app,lambda:bool(threads))
    assert threads[0]!=main_thread
    assert window.generate_button.isEnabled()
    window.generate();wait_until(app,lambda:not window.busy())
    assert window.import_active()
    assert window.project.counts(window.current_node)[1]==3
    release.set();wait_until(app,lambda:not window.import_active())
    assert len(window.project.rows(window.current_node))==4


def test_background_save_io_does_not_hold_project_lock(tmp_path,monkeypatch):
    import background_save
    p=Project.create(tmp_path,'Book');p.enable_background_save()
    entered=threading.Event();release=threading.Event();dump=background_save.json.dump
    def slow_dump(*a,**kw):
        entered.set();release.wait(5);return dump(*a,**kw)
    monkeypatch.setattr(background_save.json,'dump',slow_dump)
    p.append_text(p.data['root'],'Первый.',450)
    assert entered.wait(3)
    start=time.monotonic()
    p.append_text(p.data['root'],'Второй.',450)
    assert time.monotonic()-start<.2
    release.set();p.close_saves()
    assert len(Project.open(tmp_path).rows(p.data['root']))==2


def test_linear_split_preserves_large_text():
    text=('Глава и предложение.\n'*100000)+'  '
    parts=split_text(text,450)
    assert ''.join(parts)==text
    assert all(len(x)<=452 for x in parts)


def test_selection_survives_pagination(window):
    p=empty(window)
    from project import segment
    p.node(p.data['root'])['segments']=[segment(f'Фрагмент {i}.') for i in range(650)]
    window.refresh();click_row(window,0)
    first=window.selected_ids()[0]
    window.change_page(1);click_row(window,2,Qt.ControlModifier)
    assert len(window.selected_ids())==2 and first in window.selected_ids()
    window.change_page(-1)
    assert window.table.item(0,0).isSelected()
    assert len(window.selected_ids())==2
    window.clear_fragment_selection();assert not window.selected_ids()


def test_structure_prompt_is_nonmodal_and_can_cancel(window,app,tmp_path):
    empty(window);source=tmp_path/'headings.md'
    source.write_text('# Глава\nТекст главы.',encoding='utf-8')
    window.start_import(source)
    wait_until(app,lambda:window.import_question is not None)
    assert window.import_question.windowModality()==Qt.NonModal
    assert window.generate_button.isEnabled()
    window.cancel_import();wait_until(app,lambda:not window.import_active())
    assert not window.project.rows(window.current_node)


def test_worker_docx_analysis_keeps_headings_and_tables(tmp_path):
    from docx import Document
    from document_import import read_blocks
    from document_structure import sections_from_blocks
    doc=Document();doc.add_heading('Часть первая',1);doc.add_paragraph('Начало.')
    doc.add_paragraph('Глава 1');doc.add_table(rows=1,cols=1).cell(0,0).text='Текст таблицы'
    path=tmp_path/'chapters.docx';doc.save(path)
    progress=[]
    blocks=list(read_blocks(path,lambda *args:progress.append(args),threading.Event()))
    sections=sections_from_blocks(blocks)
    assert [s['title'] for s in sections]==['Часть первая','Глава 1']
    assert 'Текст таблицы' in sections[1]['text']
    assert progress[-1][2]>0
